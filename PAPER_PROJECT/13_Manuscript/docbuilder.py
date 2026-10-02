"""docbuilder.py — tiny content-block model rendered to .docx (python-docx) and .md from the same blocks.
Blocks: ('title', text) ('authors', text) ('h1', text) ('h2', text) ('p', text) ('bullets', [text,...]) ('eq', latex, label)
        ('fig', path, caption, width_in) ('table', header, rows, caption, col_widths_in or None) ('refs', [text,...]) ('note', text)
Inline markup in text: **bold**, *italic*, ^{sup}, _{sub} (rendered in docx runs; kept literal in markdown).
Equations are rendered to PNG with matplotlib mathtext (no TeX needed) and embedded."""
import os, re, io
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

EQ_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'equations')
os.makedirs(EQ_DIR, exist_ok=True)
TOKEN = re.compile(r'(\*\*.+?\*\*|\*.+?\*|\^\{.+?\}|_\{.+?\})')


def render_equation(latex, name):
    path = os.path.join(EQ_DIR, name + '.png')
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, f'${latex}$', fontsize=13)
    fig.savefig(path, dpi=300, bbox_inches='tight', pad_inches=0.04, transparent=False, facecolor='white')
    plt.close(fig)
    return path


def add_runs(par, text, size=None, bold=None, italic=None):
    for tok in TOKEN.split(text):
        if not tok:
            continue
        r = None
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2 and not tok.startswith('**'):
            r = par.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith('^{'):
            r = par.add_run(tok[2:-1]); r.font.superscript = True
        elif tok.startswith('_{'):
            r = par.add_run(tok[2:-1]); r.font.subscript = True
        else:
            r = par.add_run(tok)
        if size: r.font.size = Pt(size)
        if bold is not None and r.bold is None: r.bold = bold
        if italic is not None and r.italic is None: r.italic = italic
    return par


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_color); tcPr.append(shd)


def build_docx(blocks, path, base_font='Times New Roman', size=11):
    doc = Document()
    st = doc.styles['Normal']; st.font.name = base_font; st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), base_font)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.0); s.top_margin = s.bottom_margin = Inches(0.9)
    for h, sz in (('Heading 1', 14), ('Heading 2', 12)):
        doc.styles[h].font.name = base_font; doc.styles[h].font.size = Pt(sz); doc.styles[h].font.color.rgb = RGBColor(0, 0, 0)
    fig_n = tab_n = eq_n = 0
    for b in blocks:
        kind = b[0]
        if kind == 'title':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_runs(p, b[1], size=16, bold=True)
        elif kind == 'authors':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_runs(p, b[1], size=10.5, italic=True)
        elif kind == 'h1':
            doc.add_heading(b[1], level=1)
        elif kind == 'h2':
            doc.add_heading(b[1], level=2)
        elif kind == 'p':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; add_runs(p, b[1]); p.paragraph_format.space_after = Pt(6)
        elif kind == 'note':
            p = doc.add_paragraph(); add_runs(p, b[1], size=9.5, italic=True); p.paragraph_format.space_after = Pt(4)
        elif kind == 'bullets':
            for t in b[1]:
                p = doc.add_paragraph(style='List Bullet'); add_runs(p, t); p.paragraph_format.space_after = Pt(2)
        elif kind == 'eq':
            eq_n += 1
            png = render_equation(b[1], f'eq{eq_n:02d}')
            tb = doc.add_table(rows=1, cols=2); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
            c0, c1 = tb.rows[0].cells; c0.width = Inches(5.6); c1.width = Inches(0.8)
            c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            from PIL import Image
            wpx, hpx = Image.open(png).size
            c0.paragraphs[0].add_run().add_picture(png, width=Inches(min(5.4, 0.85 * wpx / 300.0)))
            c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT; c1.paragraphs[0].add_run(f'({b[2] if len(b) > 2 and b[2] else eq_n})')
        elif kind == 'fig':
            fig_n += 1
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists(b[1]): p.add_run().add_picture(b[1], width=Inches(b[3] if len(b) > 3 and b[3] else 6.4))
            else: p.add_run(f'[figure file missing: {os.path.basename(b[1])}]')
            c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; add_runs(c, f'**Fig. {fig_n}.** ' + b[2], size=9.5)
        elif kind == 'table':
            tab_n += 1
            c = doc.add_paragraph(); add_runs(c, f'**Table {tab_n}.** ' + b[3], size=9.5)
            header, rows = b[1], b[2]
            t = doc.add_table(rows=1, cols=len(header)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, h in enumerate(header):
                cell = t.rows[0].cells[i]; cell.paragraphs[0].text = ''; add_runs(cell.paragraphs[0], str(h), size=8.5, bold=True); set_cell_bg(cell, 'E7E6E6')
            for r in rows:
                cells = t.add_row().cells
                for i, v in enumerate(r):
                    cells[i].paragraphs[0].text = ''; add_runs(cells[i].paragraphs[0], str(v), size=8.5)
            if len(b) > 4 and b[4]:
                for row in t.rows:
                    for i, w in enumerate(b[4]):
                        row.cells[i].width = Inches(w)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == 'refs':
            for i, t in enumerate(b[1], 1):
                p = doc.add_paragraph(); add_runs(p, f'[{i}] ' + t, size=9.5); p.paragraph_format.space_after = Pt(2); p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.first_line_indent = Inches(-0.3)
        elif kind == 'pagebreak':
            doc.add_page_break()
    doc.save(path)


def build_md(blocks, path):
    out = []; fig_n = tab_n = eq_n = 0
    for b in blocks:
        k = b[0]
        if k == 'title': out.append(f'# {b[1]}\n')
        elif k == 'authors': out.append(f'*{b[1]}*\n')
        elif k == 'h1': out.append(f'\n## {b[1]}\n')
        elif k == 'h2': out.append(f'\n### {b[1]}\n')
        elif k in ('p', 'note'): out.append(b[1] + '\n')
        elif k == 'bullets': out.append('\n'.join(f'- {t}' for t in b[1]) + '\n')
        elif k == 'eq':
            eq_n += 1; out.append(f'\n$$ {b[1]} \\qquad ({b[2] if len(b) > 2 and b[2] else eq_n}) $$\n')
        elif k == 'fig':
            fig_n += 1; out.append(f'\n![Fig. {fig_n}]({os.path.relpath(b[1], os.path.dirname(path))})\n\n**Fig. {fig_n}.** {b[2]}\n')
        elif k == 'table':
            tab_n += 1; out.append(f'\n**Table {tab_n}.** {b[3]}\n\n| ' + ' | '.join(str(h) for h in b[1]) + ' |\n|' + '---|' * len(b[1]))
            for r in b[2]: out.append('| ' + ' | '.join(str(v) for v in r) + ' |')
            out.append('')
        elif k == 'refs':
            out.append('\n'.join(f'[{i}] {t}' for i, t in enumerate(b[1], 1)) + '\n')
        elif k == 'pagebreak': out.append('\n---\n')
    open(path, 'w').write('\n'.join(out))
