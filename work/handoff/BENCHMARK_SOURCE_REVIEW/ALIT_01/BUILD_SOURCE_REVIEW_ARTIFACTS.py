"""Build source-reading artifacts. This is NOT a thermal/mechanical solver.
All numeric reference values below are transcribed from supplied papers;
no benchmark, inversion, simulation or new-material prediction is evaluated.
"""
from pathlib import Path
import json,csv,hashlib,zipfile,re
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo

ROOT=Path('/home/user');REV=ROOT/'BENCHMARK_SOURCE_REVIEW/ALIT_01';DATE='2026-09-30'
LOCAL=REV/'REFERENCE_EXCERPTS_LOCAL';LOCAL.mkdir(exist_ok=True)
# File integrity, not scientific numerical verification.
p=ROOT/'RESEARCH_PROJECT_INPUT/LITERATURE/USER_SUPPLIED_BENCHMARKS/alit.zip';data=p.read_bytes()
blob=hashlib.sha1(b'blob '+str(len(data)).encode()+b'\0'+data).hexdigest()
m=json.loads((REV/'SOURCE_MANIFEST.json').read_text());m['actual_git_blob_sha']=blob;m['git_blob_matches_api']=(blob==m['git_blob_sha']);assert m['git_blob_matches_api'];(REV/'SOURCE_MANIFEST.json').write_text(json.dumps(m,indent=2))
pdfs=json.loads((REV/'PDF_MANIFEST.json').read_text())
for r in pdfs:
 r['inspection_status']='FULL_SOURCE_AVAILABLE — RELEVANT_FORMULATION_BCS_ICS_PARAMETERS_BENCHMARKS_INSPECTED'
 r['inspection_scope']='Selected scientific sections and key rendered equations/tables; not a comprehensive proof/admissibility audit'
 r['numerical_reproduction_status']='NOT_RUN'
 r['version_limitation']='Supplied paginated publisher-formatted copy; not independently certified latest corrected version'
(REV/'PDF_MANIFEST.json').write_text(json.dumps(pdfs,indent=2))

ops=[
 ('CV','1+tau*s','1','k','CV parent','B01 Eq5/Table1'),
 ('MCV1','1+2*tau*s','1+tau*s','k','First-order phase-lag kernel; restricted family related to MCV2 by consistent k/tau reparameterization','B01 Eq13/Table1'),
 ('MCV2','1+tau*s','2+tau*s','2*k','Effective DC differs from CV at the same unit k','B01 Eq18/Table1'),
 ('MCV3','1+3*tau*s+tau^2*s^2','1+tau*s','k','Proposed first modified source-case candidate; not final project/model approval','B01 Eq21/Table1'),
 ('MCV4','1+3*tau*s+tau^2*s^2','2+tau*s','2*k','Match total DC conductivity before interpreting mechanism controls','B01 Eq24/Table1'),
 ('MCV5','(1+tau*s)^2','2*(1+tau*s)','2*k','Equal-unit/compatible-history kernel reduces to CV with conductivity2k','B01 Eq28/Table1'),
 ('MCV6','1+3*tau*s+tau^2*s^2','1+2*tau*s','k','Distinct equal-unit numerator from MCV3; admissibility/new-material scope not audited','B01 Eq31/Table1'),
 ('GN','tau*s','1+tau*s','DC pole','TypeIII structure; k*=k/tau in the source parameterization','B01 Eq8/Table1'),
 ('MGN1','2*tau*s+tau^2*s^2','1+tau*s','DC pole','Not finite static Fourier conductance','B01 Eq34/Table1'),
 ('MGN2','tau*s+tau^2*s^2','1+2*tau*s','DC pole','Eq36 internal mismatch; independently derived equal law matchesEq37/Table1; do not silently implement a correction','B01 Eq35–37/Table1'),
 ('MGN3','2*tau*s*(1+tau*s)','(1+tau*s)^2','DC pole','Equal-unit/compatible-history kernel reduces to GN with k/2','B01 Eq40/Table1'),
 ('MGN4','tau*s+tau^2*s^2','1+3*tau*s+tau^2*s^2','DC pole','Source Table1 transcribed; complete derivation/admissibility audit not done','B01 Eq43/Table1')]
operator_records=[{'model':a,'M_s':b,'N_s':c,'DC_k_times_N_over_M':d,'note':e,'source_location':f,'origin_status':'SOURCE_FORMULA_TRANSCRIPTION + STATED_MANUAL_ALGEBRA_CHECKS','solver_status':'NOT_RUN'} for a,b,c,d,e,f in ops]
with (REV/'MODEL_OPERATORS.csv').open('w',newline='') as f:
 w=csv.DictWriter(f,fieldnames=operator_records[0]);w.writeheader();w.writerows(operator_records)
(REV/'MODEL_OPERATORS.json').write_text(json.dumps(operator_records,indent=2))

cases=[
 {'case_id':'S1_HUANG_2025_FIG5','source_id':'B01','role':'Recent source-case mathematical verification proposal','models':'CV parent; MCV3 after parent diagnostics','geometry':'1D semi-infinite longitudinal continuum x>=0','thermal_BC':'theta(0,t)=T0*H(t); dimensionless theta*(0,t*)=H(t*)','mechanical_BC':'sigma(0,t)=0; u/theta decay as x->infinity','IC_history':'Causal zero-prehistory inferred from Laplace equations; extra memory histories not fully explicitly listed in the source section; reconstruction assumption must be recorded','reported_inputs':'lambda=77.6GPa;mu=38.6GPa;alpha=1.78e-5(header unit issue);rho=8945kg/m3;cE=381J/(kgK);T0=293K','nondimensional_settings':'t*=0.06;tau*=0.04; x* range shown inFig5','method_source':'Closed decay modes/amplitudes Eqs63–66; MATLAB inversion citationBrancik1999; code/settings not included','target_quantities':'dimensionless theta(x*),sigma(x*),u(x*)','reference_availability':'Fig5 curves + analytical Laplace formula; no raw arrays/digitization performed','units_and_mapping':'Stress normalized bymu; g=gamma/(rho*cE) bysource algebra, so physical k not needed for dimensionless case. k still needed for SI conversion','critical_issue':'Record alpha as1/K dimensional interpretation and compatible initial histories before implementation; no silent correction','source_location':'B01 PDFpp8–12,Table1–2,Eqs44–66,Fig5','readiness':'SOURCE_SPECIFICATION_AVAILABLE — PLANNING/ISSUE_DISPOSITION_PENDING','computed_result':'NOT_RUN','physical_validation_scope':'None executed; source theoretical comparison, not experimental data'},
 {'case_id':'S2_BAGRI_2007_LS_CYLINDER','source_id':'B02','role':'Curved coupled common-law verification proposal','models':'LS (common Cattaneo/classical-stress law with B01CV)','geometry':'Isotropic homogeneous long hollow cylinder, plane strain,m=1,a=1,b=2','thermal_BC':'Inner q_r=f(t); outer theta=0; f(t)=1-(1+100t)*exp(-100t)','mechanical_BC':'Inner u=0; outer sigma_rr=0','IC_history':'Source Eq27:zero u,u_t,theta,theta_t,q','reported_inputs':'Dimensionless c1=1;c2=0.535;epsilon=0.02;cK=1;tau0=4. Not a measured/named copper specimen','nondimensional_settings':'Published radial profiles, Figs2–4; source-specific dimensionless time/radius','method_source':'Modified-Bessel Laplace solution Eqs43–53;Honig/Hirdes1984 inversion cited','target_quantities':'theta(r,t),sigma_rr(r,t),sigma_hoop(r,t)','reference_availability':'Published curves and closed formulas; no raw profile tables/digitization performed','units_and_mapping':'Source Eqs17/22 scaling; differs from B01. Inner outward normal is -e_r; handle flux sign','critical_issue':'Do not replace fixed inner BC with traction-free or confuse generic source parameters with B01Cu; GL/GNII are not direct unmodified B01 comparisons','source_location':'B02 PDFpp3–8,Eqs15–56,Sec4.1,Figs2–4','readiness':'SOURCE_SPECIFICATION_AVAILABLE — NUMERICAL_RECONSTRUCTION_NOT_STARTED','computed_result':'NOT_RUN','physical_validation_scope':'None executed; mathematical source case'},
 {'case_id':'S3_GORDELIY_2008_UNIFORM','source_id':'B03','role':'Heat-only classical geometry verification proposal','models':'Classical diffusion T_t=laplacianT','geometry':'Infinite exterior of two circles; centre-distance length scale;R2=0.3,R1=0.2/0.3/0.5/0.6','thermal_BC':'L1=1,L2=0 for t>0','mechanical_BC':'Not applicable to source:heat-only case','IC_history':'T(x,0)=0 temperature increment; bounded-at-infinity source problem','reported_inputs':'Source geometric ratios and prescribed temperatures; not new material coefficients','nondimensional_settings':'t=1,10,100 plussteady;A/B local polar coordinates insourceFig2','method_source':'Fourier/Bessel addition and Laplace integral;N1=N2,source BC accuracy diagnostics; direct/alternating solve variants','target_quantities':'dimensionless temperature atA/B; thermal-only','reference_availability':'Table1 has32printed values visually transcribed locally; not independently recomputed','units_and_mapping':'t=kappa_diffusivity*time/rho_centre_distance^2; source far-field boundedness, not arbitrary finite cold boundary','critical_issue':'Reported epsilon is chiefly boundary-error target, not universal comparison tolerance; printed5decimals; do not call thermal match a stress validation','source_location':'B03 PDFpp4–11,Eqs2–28,Sec7.1,Table1/Fig2','readiness':'SOURCE_VALUES_AVAILABLE — NUMERICAL_RECONSTRUCTION_NOT_STARTED','computed_result':'NOT_RUN','physical_validation_scope':'None executed; source-computed table not experiment'},
 {'case_id':'S4_GORDELIY_2008_COSINE','source_id':'B03','role':'Angular-loading thermal geometry verification proposal','models':'Classical diffusion T_t=laplacianT','geometry':'Same two-circle unbounded source geometry asS3','thermal_BC':'L1=cos(phi1),L2=0','mechanical_BC':'Not applicable to source:heat-only case','IC_history':'Zero relative-temperature initial field','reported_inputs':'Same geometric source ratios; sign-changing prescribed increment','nondimensional_settings':'t=1,10,100 plussteady;A/B sourceFig2','method_source':'Source Fourier/Bessel/Laplace formulation','target_quantities':'dimensionless temperature atA/B','reference_availability':'Table2 has32printed values visually transcribed locally; not independently recomputed','units_and_mapping':'Negative increments are intentional, not negative absoluteKelvin values','critical_issue':'Do not treat source nonmonotonic/sign-changing diffusion response as automatic evidence of non-Fourier physics','source_location':'B03 PDFp11,Sec7.2,Table2','readiness':'SOURCE_VALUES_AVAILABLE — NUMERICAL_RECONSTRUCTION_NOT_STARTED','computed_result':'NOT_RUN','physical_validation_scope':'None executed; not a coupled mechanical benchmark'}
]
(REV/'BENCHMARK_CASES.json').write_text(json.dumps(cases,indent=2))

issues=[
 {'issue_id':'Q01','source':'B01Eq36/37','observation':'Equal-parameter substitution inEq36 does not giveEq37/Table1MGN2','evidence':'Native text + renderedPDFp7 + manual eliminationEq35','disposition':'ACTION_REQUIRED beforeMGN2; no silent source/code change; CV/MCV3 scope avoids this branch','check_type':'ANALYTICALLY_CHECKED — LIMITED_SOURCE_CONSISTENCY'},
 {'issue_id':'Q02','source':'B01Table1/Eq28/40','observation':'RestrictedMCV5 andMGN3 reduce to parent kernels with changedk','evidence':'Commonfactor cancellation forcompatible causal histories','disposition':'Do not count everylabel asnew independent physics or claim first discovery; arbitrary extraICs/unequalparameters outside identity','check_type':'ANALYTICALLY_CHECKED — RESTRICTED_KERNEL_IDENTITY'},
 {'issue_id':'Q03','source':'B01Table1','observation':'MCV2/4/5 haveDC2k,notk;GN/MGNhaveDCpole','evidence':'DirectN(0)/M(0)transferlimit','disposition':'Reproduce sourceinputs faithfully; matchedstaticconductivity controls separate andlogged','check_type':'ANALYTICALLY_CHECKED — NORMALIZATION'},
 {'issue_id':'Q04','source':'B01Table2/Eq46','observation':'m/K-stylealpha header inconsistent withstress-coupling coefficientdimension','evidence':'RenderedPDFp11,stress-law dimensional requirement','disposition':'Proposedsame numeric1.78e-5 interpreted1/K; recordbeforecode; noinputaltered','check_type':'ANALYTICALLY_CHECKED — UNIT_DISPOSITION_PENDING'},
 {'issue_id':'Q05','source':'B01Eq53/Fig4','observation':'Physicalk/inversiondetails nottabulated; Fig4comparisonsettings notfullystated','evidence':'Fullsource sections/tables inspected','disposition':'DimensionlessFig5possible usingg cancellation; noSIwave claim; Ref52optional forFig4audit','check_type':'SOURCE_INSPECTED — RECONSTRUCTION_BOUNDARY'},
 {'issue_id':'Q06','source':'B01/B02','observation':'CommonCV/LSlaw butdifferentgeometry,BCs,scales;GL/GNII notdirectrecentcase','evidence':'Constitutive/energy equations andBCsections','disposition':'Separatecase configurations; no directcurve overlay bylabel','check_type':'ANALYTICALLY_CHECKED — CASE_COMPATIBILITY'},
 {'issue_id':'Q07','source':'B03','observation':'Unbounded,temperature-increment table withsource BCdiagnostic/5decimalprinting','evidence':'SourceSec2,6,7 andrenderedTables1/2','disposition':'Nofinitecoldsteadyboundarysubstitution ornegativeKelvin interpretation; referencevalues notexperimental/newresults','check_type':'SOURCE_INSPECTED — SCOPE/UNCERTAINTY'}
]
(REV/'SOURCE_ISSUES.json').write_text(json.dumps(issues,indent=2))

# Printed table transcriptions. These are input references, not outputs of a computation here.
uniform={
0.2:([0.15440,0.34789,0.43136,0.60023],[0.00606,0.06030,0.11658,0.27223]),
0.3:([0.20006,0.41523,0.50605,0.68454],[0.00847,0.07453,0.14011,0.31546]),
0.5:([0.29869,0.53339,0.62918,0.80787],[0.01528,0.10614,0.18788,0.39175]),
0.6:([0.35287,0.58714,0.68154,0.85353],[0.02031,0.12508,0.21427,0.42868])}
cosine={
0.2:([-0.08039,-0.14587,-0.16861,-0.20979],[0.00177,-0.00049,-0.01170,-0.04908]),
0.3:([-0.12365,-0.21051,-0.23908,-0.28832],[0.00298,0.00183,-0.01228,-0.05972]),
0.5:([-0.22131,-0.33428,-0.36755,-0.41880],[0.00683,0.01293,-0.00278,-0.05944]),
0.6:([-0.27677,-0.39575,-0.42881,-0.47642],[0.01005,0.02282,0.00826,-0.04870])}
printed=[]
for table,vals in [('1_uniform',uniform),('2_cosine',cosine)]:
 for radius,points in vals.items():
  for point,seq in zip(['A','B'],points):
   for t,value in zip(['1','10','100','steady'],seq):
    printed.append({'source_id':'B03','table':table,'R1':radius,'R2':0.3,'point':point,'time':t,'theta_reference_printed':f'{value:.5f}','reported_bc_eps':'1e-5' if radius==0.6 else '1e-6','origin':'SOURCE_REPORTED_COMPUTED_TABLE — NOT_EXPERIMENT — NOT_OUR_SOLVER_RESULT','transcription_check':'VISUALLY_CHECKED_AGAINST_SUPPLIED_PDF','our_result':'NOT_RUN'})
assert len(printed)==64
with (LOCAL/'B03_PRINTED_REFERENCE_VALUES.csv').open('w',newline='') as f:
 w=csv.DictWriter(f,fieldnames=printed[0]);w.writeheader();w.writerows(printed)
(LOCAL/'REFERENCE_DATA_ORIGIN.json').write_text(json.dumps({'source':'B03 suppliedPDFTables1–2,PDFpages10–11','doi':'10.1016/j.ijheatmasstransfer.2007.10.021','rows':64,'source_pdf_sha256':next(x['sha256'] for x in pdfs if x['source_id']=='B03'),'processing':'Manual/source-assisted table transcription with rendered-page visual checking; NOT numerical reproduction','precision':'5printeddecimals; sourceepsilonreferschieflytoboundarydiagnostic; acceptancecriterionnotfrozen','redistribution':'Localresearchreferenceextraction; excludedfromcheckpointpendingrightsreview'},indent=2))

wb=Workbook();wb.remove(wb.active)
def tab(name,records):
 ws=wb.create_sheet(name);cols=list(records[0]);ws.append(cols)
 for r in records:ws.append([r.get(c,'') for c in cols])
 ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
 for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='183449');cell.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=35
 for row in ws.iter_rows(min_row=2):
  for c in row:
   c.font=Font(name='Calibri',size=10);c.alignment=Alignment(wrap_text=True,vertical='top');c.fill=PatternFill('solid',fgColor='F1F4F6' if c.row%2==0 else 'FFFFFF')
   if c.value and any(x in str(c.value) for x in ['NOT_RUN','ACTION_REQUIRED','PENDING']):c.fill=PatternFill('solid',fgColor='FFF1CC')
  ws.row_dimensions[row[0].row].height=105
 for i,col in enumerate(cols,1):ws.column_dimensions[ws.cell(1,i).column_letter].width=16 if col in ['source_id','case_id','model','pages'] else 45
 t=Table(displayName=name.replace('_','')+'Index',ref=ws.dimensions);t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=False);ws.add_table(t)
 return ws
readme=[{'item':'Scope','status':'Requested source reading/compatibility screen, not fullformalPhase0A2/2 orsource-runapproval.'},
 {'item':'Access','status':'All3requestedPDFsreceived;DOIsmatched;18+11+15pages. Sourceaccessblockresolved.'},
 {'item':'Inspection','status':'Relevantfull-sourceformulation,BCs/ICs,methods,parameters,figures/tablesinspected;notcomprehensiveproofaudit.'},
 {'item':'Verification','status':'Limitedmanualsourcealgebra ANALYTICALLY_CHECKED; numericalsolver/benchmark NOT_RUN.'},
 {'item':'Physicalvalidation','status':'NOT_STARTED. Threearchivepapersprovide mathematical referencecases, notmatchednew-materialexperiments.'},
 {'item':'Referenceinputs','status':'64printedB03valuestranscribedlocally;SOURCE_COMPUTED,NOTourresults. Complete table dataexcludedfromcheckpointpendingrightsreview.'},
 {'item':'Issues','status':'MGN2Eq36mismatch;alphatableunit;kernelreductions/DCnormalization;ICandinversiondetails.'},
 {'item':'Next','status':'Proposedreference-onlypilot:dimensionlessHuangCV+MCV3. Requires explicit phase/scope/plan/budget approval. No beta/ellipseproduction.'}]
ids=[{'source_id':r['source_id'],'title':{'B01':'Huang2025 generalized heat conduction/thermoelasticity','B02':'Bagri/Eslami2007 cylinders/spheres','B03':'Gordeliy2008 two circular cavities'}[r['source_id']],'doi':r['doi_matches'][0],'pages':r['pages'],'file':r['filename'],'sha256':r['sha256'],'source_access':'RESOLVED','scientific_reading':r['inspection_status'],'simulation':'NOT_RUN'} for r in pdfs]
tab('Readme',readme);tab('Source_PDFs',ids);tab('Case_specs',cases);tab('Law_operators',operator_records);tab('Source_issues',issues)
wb.save(REV/'SOURCE_BENCHMARK_READINESS.xlsx')

# Update only current control/status files; earlier scientific reports/inputs remain immutable history.
def append(path,text):
 with path.open('a') as f:f.write('\n'+text+'\n')
append(ROOT/'DECISION_LOG.md','''## D006 — alit full sources received / read\nDate:30 September2026. User supplied alit.zip containing the exact3requestedpapers. Former accessblockersresolved; relevantsourceequations/BCs/ICs/material/casesinspected. Numericalverification/physicalvalidation remainunexecuted. No finalmodel/material/journalapproval inferred from upload.\n\n## D007 — Proposed reference route / source issues\nRecommend firstB01Figure5CVparent+MCV3sourcecase, thenB02LScylinder andB03heat-onlytables separately. Limitedmanualoperator/boundary/common-law checks recorded. MGN2Eq36mismatch andalphatableunit requiredocumenteddisposition;originalsourceunchanged. Match totalstaticconductivity in separatemechanismcontrols, notbysilentlychangingpublishedreferenceinputs. No sourcecorrectionimplemented, no solveroutputaffected.\n\n## D008 — Execution hold\nA possiblePhase0A2reference-onlyfeasibilitypilot hasproposedcap0.5CPU-hour/noGPU/50MB; this isnotobservedruntime/productionapproval. Needexplicitnextphase/scopeauthorizationandpre-runplan/criteria. No newmaterial/ellipsecavityproduction authorized.''')
append(ROOT/'CHANGELOG.md','''## 2026-09-30 — alit source-review checkpoint\n- Preservedalit.zip/3PDFswithhashes;matchedDOIs,44documentpages.\n- Read relevantfull-source equations/cases; sourceaccessB01–B03nowresolved. Priorblockedpreviewreportsretainedashistory.\n- Addedcaseblueprints,operatorindex,manualanalyticalsourcechecks andprivateprintedtablereferences; not newsolverresults.\n- RecordedMGN2internalmismatch,alpha-unitdisposition and restrictedkernelreductions/DCnormalization. No originalPDF/model/code modified.\n- Updatedlateststate/currentstatuses; proposedpilotnotapproved/run.\n- Recoverycheckpoint excludesprotectedinputs/fulltext/images/completetableextractions.''')
(ROOT/'VALIDATION_STATUS.md').write_text('''# Current verification/physical-validation status\n\nDate:30 September2026. MASTER_PROMPTv2.1. Lateststate:PROJECT_STATE_FULL_SOURCE_REVIEW_01.md.\n\nNewresearchdirection/model/material/evidenceroute:NOT_APPROVED.\n\n## Verification\nNumericalimplementation/benchmark/convergence:NOT_STARTED / NOT_RUN.\nLimitedsourcealgebra/units/boundaryamplitude/common-lawchecks:ANALYTICALLY_CHECKED with stated parameter/historyconditions, not wholeprojectverification.\nSourceaccessB01–B03:RESOLVED, relevantfulltextequations/BCs/casesinspected. Source-caseplansavailable; additionalhistory/unit/inversion/reference-uncertaintydetails needpre-runrecord.\nMGN2sourceinternalmismatch:ACTION_REQUIRED ifbranchselected; originalsourceunchanged.\n\n## Physicalvalidation\nNOT_STARTED. Currentuploaded3paperscontain theoretical/computationalcases, notmatchednew-materialT(t)/stressmeasurements. ExistingB04/B05material-property evidence is not a new-model prediction comparison. Forproposednon-Fourierbeta-Ga2O3transientstress/temperaturepredictions, matchedsuitablephysicalevidence not established:APPLICABLE — EVIDENCE_UNAVAILABLE (provisional boundedassessment,notproofnoneexists). NOT_APPLICABLE not justified.\n\nSourcevaluesare inputs/referencevalues, not ourcomputedresults. Publishedvalidationheadings/literaturecomparison are notautomaticallyphysicalvalidation. No quantitativeagreement/convergence/validatedscope claimed.\n''')
append(ROOT/'COMPUTE_STATUS.md','''## alit source-review update\nScientificsolver/pilot/production runsremain0. Executeddocumentoperations:ZIP/PDFreading,textextraction,sourcepagerendering,metadata/tabletranscription andworkbook/checkpointcreation. Manualanalyticalsourcechecksare notnumericalsolverruns. Proposedpilotcap0.5CPU-hour/noGPU/50MB isnot anobservedruntime orapproval; method/runtimeestimates not profiled.''')
append(ROOT/'REPRODUCIBILITY_STATUS.md','''## alit source-review update\nAll3requestedPDFsavailable atrecordedpaths/hashes; source-case equation/conditionnotes andlocalprintedtableinputsavailable. Authorcode/rawcoupledprofilearrays notsupplied. Numericalreference/convergence/principalresultreproduction NOT_RUN. Newnotescheckpoint deliberatelyexcludes protectedsourceinputs/fulltexts/images/completetabletranscriptions. No offline-complete/finalresearcharchive claimed.''')
append(ROOT/'RESEARCH_PROFILE.md','''## alit source update\nRequestedrecent-model/coupled/thermalreferencePDFsreceived andread. SourcebenchmarkmaterialCu/genericnondimensionalcases do not select a newresearchmaterial; beta-Ga2O3/ellipse/currentmodelbranch remain provisional. No experimentalnew-materialdataset, authorcode orcomputepreferences supplied inalit.zip. Sourcesare nowavailablefor a scopedreference-caseplan; executionauthorizationstillpending.''')
append(ROOT/'REJECTED_IDEAS.md','''## Source-review holds\n- Do not call all12equal-unitlabels independentnewphysics: restrictedMCV5/MGN3simplify;MCV1/2related under consistentrescaling.\n- Do not implement ambiguousMGN2Eq36silently. Branchheldpending documentedsourceconsistencydisposition,notwholesalepaperrejection.\n- Do not treat sameunitk acrossall sourcebranches as equal measuredtotalDCconductivity.\n- Do not treat B02GL/GNIIcurves as automaticB01unmodifiedmodelvalidation.\n- Do not usefinitecoldsteadydomain as identicalB03unboundedcase or negativeincrements as negativeKelvin.\nNo user-approveddirection wasrejected/replaced.''')
(ROOT/'CURRENT_PROJECT_STATE.md').write_text('''# Current state pointer\n\nRead MASTER_PROMPT.md (v2.1), then PROJECT_STATE_FULL_SOURCE_REVIEW_01.md and rootdecision/statusfiles. Three requestedfullsources are available; reference/spec/analyticalsourcechecks reviewed. No scientificsolver/pilot/production run or finaldirectionapproval. STOP before execution until explicitnext scoped-phaseapproval.\n''')
# PDF/case/index artifacts checked; do not assert any scientific numerical outcome.
assert len(pdfs)==3 and sum(x['pages'] for x in pdfs)==44
assert len(load_workbook(REV/'SOURCE_BENCHMARK_READINESS.xlsx',read_only=True).sheetnames)==5
assert all(x['computed_result']=='NOT_RUN' for x in cases)
print('SOURCE_ARTIFACT_BUILD_RUN_SUCCESS — DOCUMENT/REFERENCE_PROCESSING_ONLY')
print('B03_SOURCE_REFERENCE_ROWS',len(printed))
print('MODEL_OPERATOR_ROWS',len(operator_records))
print('GIT_BLOB_MATCHES',m['git_blob_matches_api'])
print('Scientificsolver/benchmark/pilot/production NOT_RUN')
